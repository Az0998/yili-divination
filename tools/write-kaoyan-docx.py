# -*- coding: utf-8 -*-
from pathlib import Path
import json
import sys

try:
    from docx import Document
except ImportError:
    sys.stderr.write("python-docx missing\n")
    sys.exit(1)

root = Path(__file__).resolve().parents[1]
out = root / "out"
md_path = out / "ZhangSenjie-20260918-ChenHour-Kaoyan.md"
json_path = out / "ZhangSenjie-20260918-ChenHour-Kaoyan.json"
docx_path = out / "ZhangSenjie-20260918-ChenHour-Kaoyan.docx"

doc = Document()
doc.add_heading("张森捷 2028 跨考规划（GitHub 备用辰时三筮）", 0)

if json_path.exists():
    data = json.loads(json_path.read_text(encoding="utf-8"))
    p = doc.add_paragraph()
    p.add_run("占时：").bold = True
    p.add_run(data.get("castTimeRecorded", ""))
    doc.add_paragraph(data.get("note", ""))
    doc.add_heading("大学层 7 优胜", 1)
    for row in data.get("universitiesTop7") or []:
        doc.add_paragraph(
            f"{row.get('rank')}. {row.get('label')}  {row.get('bengua')}之{row.get('biangua')}  {row.get('score')}分  {row.get('verdict')}"
        )
    doc.add_heading("专业层 7 优胜", 1)
    for row in data.get("majorsTop7") or []:
        doc.add_paragraph(
            f"{row.get('rank')}. {row.get('label')}  {row.get('bengua')}之{row.get('biangua')}  {row.get('score')}分  {row.get('verdict')}"
        )
    doc.add_heading("冲 / 稳组合", 1)
    for line in data.get("stretchVsSafe") or []:
        doc.add_paragraph(line)
    doc.add_heading("导师层", 1)
    doc.add_paragraph("不编造导师姓名。请用以上学校当年公开招生名单核对。")
elif md_path.exists():
    for line in md_path.read_text(encoding="utf-8").splitlines():
        if line.startswith("# "):
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:], 1)
        else:
            doc.add_paragraph(line)
else:
    sys.stderr.write("missing json/md\n")
    sys.exit(1)

doc.save(docx_path)
print(str(docx_path))
